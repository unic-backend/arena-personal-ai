"""Lecture des documents du propriétaire.

Les fichiers d'essai sont de **vrais** PDF et de **vrais** `.docx`, fabriqués
dans le test : un lecteur de PDF vérifié sur une chaîne de caractères ne prouve
rien.
"""
import shutil

import pytest

import tools.documents.reader as reader
from tools.documents.reader import Passage, lire_document

pytest.importorskip("pypdf", reason="pypdf n'est pas installe.")
pytest.importorskip("docx", reason="python-docx n'est pas installe.")
pytest.importorskip("openpyxl", reason="openpyxl n'est pas installe.")
pytest.importorskip("pptx", reason="python-pptx n'est pas installe.")
pytest.importorskip("pypdfium2", reason="pypdfium2 n'est pas installe.")
pytest.importorskip("pytesseract", reason="pytesseract n'est pas installe.")


def fabriquer_pdf(pages: list[str]) -> bytes:
    """Construit un PDF valide, une page par entrée, sans dépendance d'écriture."""
    flux = []
    for texte in pages:
        echappe = texte.replace("\\\\", r"\\\\").replace("(", r"\(").replace(")", r"\)")
        flux.append(f"BT /F1 12 Tf 72 720 Td ({echappe}) Tj ET".encode("latin-1"))

    n = len(pages)
    objets = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{3 + 2 * i} 0 R' for i in range(n))}] "
        f"/Count {n} >>".encode(),
    ]
    for i, f in enumerate(flux):
        objets.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {3 + 2 * n} 0 R >> >> >>".encode()
        )
        objets.append(b"<< /Length " + str(len(f)).encode() + b" >>\nstream\n" + f + b"\nendstream")
    objets.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    sortie, positions = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objets, 1):
        positions.append(len(sortie))
        sortie += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    debut = len(sortie)
    sortie += f"xref\n0 {len(objets) + 1}\n0000000000 65535 f \n".encode()
    for p in positions:
        sortie += f"{p:010d} 00000 n \n".encode()
    sortie += (f"trailer\n<< /Size {len(objets) + 1} /Root 1 0 R >>\n"
               f"startxref\n{debut}\n%%EOF\n").encode()
    return bytes(sortie)


@pytest.fixture
def devis_pdf(tmp_path):
    chemin = tmp_path / "devis_2026_041.pdf"
    chemin.write_bytes(fabriquer_pdf([
        "Devis numero 2026-041 - cloison BA13 salon",
        "Total HT 450000 FCFA - validite 30 jours",
    ]))
    return chemin


@pytest.fixture
def plan_scanne_pdf(tmp_path):
    """Un vrai « scan » : le texte est une IMAGE, sans aucune couche texte.

    Pleine page a une resolution proche de 300 DPI — un scan reel, pas une
    vignette. Mesure du 30/08/2026 : une image plus petite ou un rendu moins
    resolu rend le texte illisible pour tesseract, meme si `pypdf` confirme
    bien l'absence de couche texte.
    """
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    image = Image.new("RGB", (2480, 3508), "white")  # A4 a ~300 DPI
    dessin = ImageDraw.Draw(image)
    dessin.text((150, 200), "Devis numero 2026-118 Chantier Almadies", fill="black")
    dessin.text((150, 300), "Cloison BA13 20.00 m2 Total HT 450000 FCFA", fill="black")
    image_chemin = tmp_path / "page_scan.png"
    image.save(image_chemin)

    chemin = tmp_path / "plan_scanne.pdf"
    c = canvas.Canvas(str(chemin), pagesize=A4)
    c.drawImage(str(image_chemin), 0, 0, width=A4[0], height=A4[1])
    c.save()
    return chemin


@pytest.fixture
def facture_docx(tmp_path):
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("Facture 2026-118 - UniC Plaquiste")
    tableau = document.add_table(rows=2, cols=3)
    tableau.rows[0].cells[0].text = "Designation"
    tableau.rows[0].cells[1].text = "Quantite"
    tableau.rows[0].cells[2].text = "Prix"
    tableau.rows[1].cells[0].text = "Plaque BA13"
    tableau.rows[1].cells[1].text = "40"
    tableau.rows[1].cells[2].text = "180000"

    chemin = tmp_path / "facture_2026_118.docx"
    document.save(str(chemin))
    return chemin


@pytest.fixture
def metre_xlsx(tmp_path):
    from openpyxl import Workbook

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Metre"
    feuille.append(["Designation", "Quantite", "Prix"])
    feuille.append(["Plaque BA13", 40, 180000])

    autre = classeur.create_sheet("Notes")
    autre.append(["Chantier Almadies"])

    chemin = tmp_path / "metre_2026_118.xlsx"
    classeur.save(str(chemin))
    return chemin


@pytest.fixture
def argumentaire_pptx(tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    diapo1 = presentation.slides.add_slide(presentation.slide_layouts[1])
    diapo1.shapes.title.text = "UniC Plaquiste"
    diapo1.placeholders[1].text = "Chantier Almadies - cloison BA13"

    diapo2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    tableau_forme = diapo2.shapes.add_table(2, 2, 0, 0, 4000000, 800000)
    tableau = tableau_forme.table
    tableau.cell(0, 0).text = "Designation"
    tableau.cell(0, 1).text = "Prix"
    tableau.cell(1, 0).text = "Plaque BA13"
    tableau.cell(1, 1).text = "180000"

    chemin = tmp_path / "argumentaire_2026_118.pptx"
    presentation.save(str(chemin))
    return chemin


# --- PDF -----------------------------------------------------------------------

def test_un_pdf_est_lu_page_par_page(devis_pdf):
    document = lire_document(devis_pdf)

    assert document.lu
    assert len(document.passages) == 2
    assert "cloison BA13" in document.passages[0].texte
    assert "450000" in document.passages[1].texte


def test_chaque_passage_d_un_pdf_garde_son_numero_de_page(devis_pdf):
    """Sans la page, une réponse ne peut pas citer précisément sa source."""
    document = lire_document(devis_pdf)

    assert [p.page for p in document.passages] == [1, 2]
    assert document.passages[1].source == "devis_2026_041.pdf, page 2"


def test_un_pdf_sans_texte_est_signale_pas_invente(tmp_path):
    """Un PDF scanné ne contient que des images : le dire vaut mieux que deviner."""
    chemin = tmp_path / "scan.pdf"
    chemin.write_bytes(fabriquer_pdf([""]))

    document = lire_document(chemin)

    assert document.statut == "VIDE"
    assert "scanne" in document.raison
    assert document.texte == ""


def test_un_pdf_corrompu_ne_fait_pas_tomber_la_lecture(tmp_path):
    chemin = tmp_path / "casse.pdf"
    chemin.write_bytes(b"%PDF-1.4\nceci n'est pas un PDF valide")

    document = lire_document(chemin)

    assert document.statut in {"ECHEC", "VIDE"}
    assert document.raison


# --- OCR (PDF scanné) -----------------------------------------------------------

@pytest.fixture
def tesseract_disponible():
    """Le binaire reel, pas seulement les paquets Python — comme `gitleaks` ailleurs."""
    if shutil.which("tesseract") is None:
        pytest.skip("le binaire tesseract n'est pas installe sur cette machine.")


@pytest.mark.integration
def test_une_page_scannee_est_lue_par_ocr(tesseract_disponible, plan_scanne_pdf):
    """Bout en bout, avec le vrai moteur — pas un moteur simule."""
    document = lire_document(plan_scanne_pdf)

    assert document.lu
    assert "2026-118" in document.texte
    assert "Almadies" in document.texte or "Amadies" in document.texte  # police de test grossiere
    assert document.passages[0].via_ocr is True
    assert document.passages[0].source == "plan_scanne.pdf, page 1 (OCR)"
    assert document.resume()["passages_ocr"] == 1


def test_une_page_sans_texte_appelle_l_ocr(monkeypatch, tmp_path):
    """Sans dependre du vrai tesseract : verifie que le relais est bien branche."""
    monkeypatch.setattr(reader, "_ocr_page", lambda chemin, index: "Texte devine par OCR")
    chemin = tmp_path / "scan.pdf"
    chemin.write_bytes(fabriquer_pdf([""]))

    document = lire_document(chemin)

    assert document.lu
    assert "Texte devine par OCR" in document.texte
    assert document.passages[0].via_ocr is True


def test_une_page_avec_texte_n_appelle_pas_l_ocr(monkeypatch, devis_pdf):
    """Une page deja lisible ne doit jamais repasser par l'OCR — coute cher pour rien."""
    def _echoue_si_appele(chemin, index):
        raise AssertionError("l'OCR a ete appele sur une page qui avait deja du texte")
    monkeypatch.setattr(reader, "_ocr_page", _echoue_si_appele)

    document = lire_document(devis_pdf)

    assert document.lu
    assert all(not p.via_ocr for p in document.passages)


def test_ocr_qui_ne_trouve_rien_reste_vide(monkeypatch, tmp_path):
    """Tesseract absent ou une page vraiment blanche : `_ocr_page` rend "" — jamais un crash."""
    monkeypatch.setattr(reader, "_ocr_page", lambda chemin, index: "")
    chemin = tmp_path / "scan.pdf"
    chemin.write_bytes(fabriquer_pdf([""]))

    document = lire_document(chemin)

    assert document.statut == "VIDE"


# --- Word ----------------------------------------------------------------------

def test_un_docx_est_lu(facture_docx):
    document = lire_document(facture_docx)

    assert document.lu
    assert "Facture 2026-118" in document.texte


def test_le_contenu_des_tableaux_est_conserve(facture_docx):
    """Un devis vit souvent dans un tableau : l'ignorer viderait le document."""
    document = lire_document(facture_docx)

    assert "Plaque BA13" in document.texte
    assert "180000" in document.texte
    assert "Designation | Quantite | Prix" in document.texte


# --- Excel ---------------------------------------------------------------------

def test_un_xlsx_est_lu_feuille_par_feuille(metre_xlsx):
    document = lire_document(metre_xlsx)

    assert document.lu
    assert len(document.passages) == 2
    assert "Plaque BA13" in document.passages[0].texte
    assert "Almadies" in document.passages[1].texte


def test_le_nom_de_la_feuille_xlsx_est_dans_la_provenance(metre_xlsx):
    document = lire_document(metre_xlsx)

    assert document.passages[0].source == "metre_2026_118.xlsx (Metre)"


def test_les_lignes_xlsx_gardent_leurs_colonnes(metre_xlsx):
    document = lire_document(metre_xlsx)

    assert "Designation | Quantite | Prix" in document.texte
    assert "Plaque BA13 | 40 | 180000" in document.texte


def test_une_feuille_xlsx_vide_ne_produit_aucun_passage(tmp_path):
    from openpyxl import Workbook

    classeur = Workbook()
    classeur.active.title = "Vide"
    chemin = tmp_path / "vide.xlsx"
    classeur.save(str(chemin))

    document = lire_document(chemin)

    assert document.statut == "VIDE"


def test_un_xlsx_corrompu_ne_fait_pas_tomber_la_lecture(tmp_path):
    chemin = tmp_path / "casse.xlsx"
    chemin.write_bytes(b"ceci n'est pas un classeur Excel")

    document = lire_document(chemin)

    assert document.statut == "ECHEC"


# --- PowerPoint ------------------------------------------------------------------

def test_un_pptx_est_lu_diapositive_par_diapositive(argumentaire_pptx):
    document = lire_document(argumentaire_pptx)

    assert document.lu
    assert len(document.passages) == 2
    assert "UniC Plaquiste" in document.passages[0].texte
    assert "Almadies" in document.passages[0].texte


def test_le_tableau_d_une_diapositive_est_conserve(argumentaire_pptx):
    document = lire_document(argumentaire_pptx)

    assert "Designation | Prix" in document.texte
    assert "Plaque BA13 | 180000" in document.texte


def test_chaque_passage_pptx_garde_son_numero_de_diapositive(argumentaire_pptx):
    document = lire_document(argumentaire_pptx)

    assert [p.page for p in document.passages] == [1, 2]
    assert document.passages[0].source == "argumentaire_2026_118.pptx, page 1"


def test_un_pptx_corrompu_ne_fait_pas_tomber_la_lecture(tmp_path):
    chemin = tmp_path / "casse.pptx"
    chemin.write_bytes(b"ceci n'est pas une presentation")

    document = lire_document(chemin)

    assert document.statut == "ECHEC"


# --- Texte ---------------------------------------------------------------------

@pytest.mark.parametrize("extension", [".txt", ".md", ".csv"])
def test_un_fichier_texte_est_lu(tmp_path, extension):
    chemin = tmp_path / f"note{extension}"
    chemin.write_text("Chantier Almadies\nCloison 12 m2", encoding="utf-8")

    document = lire_document(chemin)

    assert document.lu
    assert "Almadies" in document.texte


def test_un_encodage_inattendu_ne_fait_pas_echouer_la_lecture(tmp_path):
    chemin = tmp_path / "ancien.txt"
    chemin.write_bytes("Réunion à Dakar".encode("latin-1"))

    document = lire_document(chemin)

    assert document.lu
    assert "union" in document.texte


def test_les_lignes_vides_en_trop_sont_reduites(tmp_path):
    chemin = tmp_path / "aere.txt"
    chemin.write_text("Titre\n\n\n\n\nContenu", encoding="utf-8")

    assert lire_document(chemin).texte == "Titre\n\nContenu"


# --- Refus ---------------------------------------------------------------------

@pytest.mark.parametrize("nom", ["photo.jpg", "archive.zip", "video.mp4", "tableur.ods"])
def test_un_format_non_pris_en_charge_est_refuse_explicitement(tmp_path, nom):
    chemin = tmp_path / nom
    chemin.write_bytes(b"contenu quelconque")

    document = lire_document(chemin)

    assert document.statut == "NON_PRIS_EN_CHARGE"
    assert "formats lus" in document.raison


def test_un_fichier_absent_est_signale(tmp_path):
    document = lire_document(tmp_path / "jamais_vu.pdf")

    assert document.statut == "ECHEC"
    assert "introuvable" in document.raison


def test_un_fichier_trop_volumineux_est_refuse(tmp_path):
    chemin = tmp_path / "enorme.txt"
    chemin.write_bytes(b"x" * 5000)

    document = lire_document(chemin, taille_max=1000)

    assert document.statut == "ECHEC"
    assert "volumineux" in document.raison


def test_un_fichier_vide_est_signale(tmp_path):
    chemin = tmp_path / "vide.txt"
    chemin.write_text("", encoding="utf-8")

    assert lire_document(chemin).statut == "VIDE"


# --- Résumé --------------------------------------------------------------------

def test_le_resume_ne_contient_pas_le_texte(devis_pdf):
    """Il part dans les journaux : y mettre le contenu d'une facture serait une fuite."""
    resume = lire_document(devis_pdf).resume()

    assert resume["statut"] == "LU"
    assert resume["passages"] == 2
    assert resume["caracteres"] > 0
    assert "450000" not in str(resume)


def test_la_provenance_sans_page_reste_le_nom_du_fichier():
    assert Passage(texte="x", fichier="note.txt").source == "note.txt"
