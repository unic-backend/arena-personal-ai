from pptx import Presentation

from core.actions.resultat import Statut
from core.connectors.presentation import ConnecteurPresentation


def test_genere_un_pptx_reel_et_editable(tmp_path):
    connecteur = ConnecteurPresentation(dossier=tmp_path)
    resultat = connecteur._executer(
        connecteur.capacites()["generer"],
        plan={"titre": "Projet Arena", "theme": "sombre", "slides": [
            {"titre": "Vision", "puces": ["Assistant autonome", "Local-first"]},
            {"titre": "Étapes", "puces": ["Mesurer", "Construire", "Vérifier"]},
        ]},
    )
    assert resultat.statut == Statut.SUCCES
    fichier = next(tmp_path.glob("*.pptx"))
    relu = Presentation(fichier)
    assert len(relu.slides) == 2
    assert relu.slides[0].shapes[0].text == "Vision"
    assert resultat.detail["theme"] == "sombre"


def test_refuse_un_plan_sans_slides(tmp_path):
    connecteur = ConnecteurPresentation(dossier=tmp_path)
    resultat = connecteur._executer(
        connecteur.capacites()["generer"], plan={"titre": "Vide", "slides": []})
    assert resultat.statut == Statut.ECHEC
    assert list(tmp_path.glob("*.pptx")) == []


def test_refuse_un_json_invalide(tmp_path):
    connecteur = ConnecteurPresentation(dossier=tmp_path)
    resultat = connecteur._executer(connecteur.capacites()["generer"], plan="{")
    assert resultat.statut == Statut.ECHEC
