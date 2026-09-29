"""Word -> PDF sans LibreOffice (DEC-0173) : le PDF ecrit est relu, et rien
du document n'est interprete comme du code."""
from dataclasses import replace
from pathlib import Path

import pytest

from core.connectors.file_conversion import ConnecteurFileConversion
from core.production.conversion import registre, word_pdf
from core.production.conversion.moteurs import MoteurEchec


def _word(tmp_path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Compte rendu de chantier", level=1)
    paragraphe = document.add_paragraph("Travaux ")
    paragraphe.add_run("urgents").bold = True
    document.add_paragraph("cloisons BA13", style="List Bullet")
    document.add_paragraph("faux plafond", style="List Bullet")
    document.add_paragraph("Commander", style="List Number")
    tableau = document.add_table(rows=2, cols=2)
    tableau.cell(0, 0).text, tableau.cell(0, 1).text = "Poste", "Total"
    tableau.cell(1, 0).text, tableau.cell(1, 1).text = "BA13", "180 000"
    document.add_paragraph("<img src='http://x.test/espion.png'> & <script>alert(1)</script>")
    chemin = tmp_path / "chantier.docx"
    document.save(chemin)
    return chemin


def test_la_structure_du_word_devient_du_html(tmp_path):
    page = word_pdf.docx_vers_html(_word(tmp_path))

    assert "<h1>Compte rendu de chantier</h1>" in page
    assert "<p>Travaux <strong>urgents</strong></p>" in page
    assert "<ul><li>cloisons BA13</li><li>faux plafond</li></ul><ol><li>Commander</li></ol>" in page
    assert "<tr><td>Poste</td><td>Total</td></tr>" in page


def test_rien_du_word_n_est_interprete_comme_du_code(tmp_path):
    page = word_pdf.docx_vers_html(_word(tmp_path))

    assert "<img" not in page and "<script" not in page
    assert "&lt;img src=&#x27;http://x.test/espion.png&#x27;&gt; &amp; &lt;script&gt;" in page


def test_aucune_ressource_n_est_chargee_pendant_le_rendu():
    with pytest.raises(ValueError, match="refusee"):
        word_pdf._refuser_toute_url("http://x.test/espion.png")


def test_sans_libreoffice_un_word_devient_un_pdf(tmp_path):
    from pypdf import PdfReader

    originales = registre.MATRICE[("docx", "pdf")]
    registre.MATRICE[("docx", "pdf")] = [
        replace(e, disponible=lambda: (False, "LibreOffice absent")) if e.moteur_id == "libreoffice" else e
        for e in originales
    ]
    try:
        resultat = ConnecteurFileConversion(dossier=tmp_path / "sorties")._convertir_un_fichier(
            str(_word(tmp_path)), "pdf")
    finally:
        registre.MATRICE[("docx", "pdf")] = originales

    assert resultat.statut.value == "SUCCESS", resultat.message
    assert resultat.detail["moteur"] == "python-docx+weasyprint"
    assert "polices, marges, images" in resultat.detail["limites_qualite"]
    texte = " ".join(page.extract_text() for page in PdfReader(resultat.preuve).pages)
    for attendu in ("Compte rendu de chantier", "urgents", "cloisons BA13", "180 000", "<img src="):
        assert attendu in texte


def test_libreoffice_garde_la_priorite_quand_il_est_la():
    assert [m.moteur_id for m in registre.moteurs_pour("docx", "pdf")] == [
        "libreoffice", "python-docx+weasyprint"]


def test_un_word_vide_ou_corrompu_echoue_en_le_disant(tmp_path):
    from docx import Document

    vide = tmp_path / "vide.docx"
    Document().save(vide)
    with pytest.raises(MoteurEchec, match="aucun texte"):
        word_pdf.docx_vers_html(vide)

    casse = tmp_path / "casse.docx"
    casse.write_bytes(b"PK\x03\x04 pas un vrai docx")
    with pytest.raises(MoteurEchec, match="illisible"):
        word_pdf.docx_vers_html(casse)


@pytest.mark.parametrize("style, niveau", [
    ("Heading 1", 1), ("Heading 3", 3), ("Titre 2", 2), ("Title", 1), ("Normal", None), ("Heading 9", 6),
])
def test_le_niveau_de_titre_se_lit_dans_le_style(style, niveau):
    assert word_pdf._niveau_de_titre(style) == niveau


# --- Excel -> PDF sans LibreOffice (DEC-0174) --------------------------------------

def _classeur(tmp_path: Path) -> Path:
    from datetime import date

    from openpyxl import Workbook

    classeur = Workbook()
    ventes = classeur.active
    ventes.title = "Ventes <T3>"
    ventes.append(["Client", "Montant", "Date", "Note"])
    ventes.append(["Diallo", 1250.5, date(2026, 9, 29), "=HYPERLINK(\"http://x.test\")"])
    ventes.append(["Ndiaye", 40, None, "<img src='http://x.test/e.png'>"])
    classeur.create_sheet("Vide")
    classeur.create_sheet("Stock").append(["BA13", 120])
    chemin = tmp_path / "ventes.xlsx"
    classeur.save(chemin)
    return chemin


def test_chaque_feuille_non_vide_devient_un_tableau(tmp_path):
    from core.production.conversion import tableurs

    page = tableurs.xlsx_vers_html(_classeur(tmp_path))

    assert "<h2>Ventes &lt;T3&gt;</h2>" in page and "<h2>Stock</h2>" in page
    assert "<h2>Vide</h2>" not in page, "une feuille vide n'est pas un tableau"
    assert "<tr><th>Client</th><th>Montant</th><th>Date</th><th>Note</th></tr>" in page
    assert '<td class="nombre">1250,5</td><td>29/09/2026</td>' in page
    assert "<img" not in page and "&lt;img src=&#x27;http://x.test/e.png&#x27;&gt;" in page
    assert "<th>BA13</th>" not in page, "une ligne seule avec un nombre n'est pas un en-tete"


def test_sans_libreoffice_un_classeur_devient_un_pdf(tmp_path):
    from pypdf import PdfReader

    originales = registre.MATRICE[("xlsx", "pdf")]
    registre.MATRICE[("xlsx", "pdf")] = [
        replace(e, disponible=lambda: (False, "LibreOffice absent")) if e.moteur_id == "libreoffice" else e
        for e in originales
    ]
    try:
        resultat = ConnecteurFileConversion(dossier=tmp_path / "sorties")._convertir_un_fichier(
            str(_classeur(tmp_path)), "pdf")
    finally:
        registre.MATRICE[("xlsx", "pdf")] = originales

    assert resultat.statut.value == "SUCCESS", resultat.message
    assert resultat.detail["moteur"] == "openpyxl+weasyprint"
    lecteur = PdfReader(resultat.preuve)
    texte = " ".join(page.extract_text() for page in lecteur.pages)
    for attendu in ("Ventes <T3>", "Diallo", "1250,5", "29/09/2026", "Stock", "BA13"):
        assert attendu in texte
    largeur, hauteur = lecteur.pages[0].mediabox.width, lecteur.pages[0].mediabox.height
    assert largeur > hauteur, "paysage : un tableau est plus large que haut"


def test_un_classeur_trop_grand_est_refuse_au_lieu_d_etre_tronque(monkeypatch, tmp_path):
    from core.production.conversion import tableurs

    monkeypatch.setattr(tableurs, "LIGNES_PDF_MAX", 3)
    with pytest.raises(MoteurEchec, match="convertis ce classeur en CSV"):
        tableurs.xlsx_vers_html(_classeur(tmp_path))


def test_un_classeur_vide_ne_donne_pas_de_pdf(tmp_path):
    from openpyxl import Workbook

    from core.production.conversion import tableurs

    chemin = tmp_path / "vide.xlsx"
    Workbook().save(chemin)
    with pytest.raises(MoteurEchec, match="aucune donnee"):
        tableurs.xlsx_vers_html(chemin)


def test_libreoffice_garde_la_priorite_pour_le_tableur():
    assert [m.moteur_id for m in registre.moteurs_pour("xlsx", "pdf")] == [
        "libreoffice", "openpyxl+weasyprint"]
