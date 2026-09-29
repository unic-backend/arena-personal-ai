"""PDF / Word / Excel / PowerPoint -> texte (DEC-0170) : le texte ecrit est
relu, et chaque passage d'un document a plusieurs parties dit d'ou il vient."""
from pathlib import Path

import pytest

from core.connectors.file_conversion import ConnecteurFileConversion
from core.production.conversion.registre import moteurs_pour


@pytest.fixture()
def connecteur(tmp_path: Path) -> ConnecteurFileConversion:
    return ConnecteurFileConversion(dossier=tmp_path / "sorties")


def _texte(connecteur, source: Path) -> str:
    resultat = connecteur._convertir_un_fichier(str(source), "txt")
    assert resultat.statut.value == "SUCCESS", resultat.message
    assert resultat.detail["moteur"] == "lecteur-documents"
    assert "OCR" in resultat.detail["limites_qualite"]
    return Path(resultat.preuve).read_text(encoding="utf-8")


def test_un_pdf_donne_son_texte_page_par_page(connecteur, tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    chemin = tmp_path / "rapport.pdf"
    page = canvas.Canvas(str(chemin), pagesize=A4)
    page.drawString(72, 750, "Chantier Almadies : cloisons posees.")
    page.showPage()
    page.drawString(72, 750, "Reste a faire : faux plafond.")
    page.save()

    texte = _texte(connecteur, chemin)

    assert "--- rapport.pdf, page 1 ---\nChantier Almadies : cloisons posees." in texte
    assert "--- rapport.pdf, page 2 ---\nReste a faire : faux plafond." in texte


def test_un_word_d_un_seul_tenant_n_a_pas_d_en_tete(connecteur, tmp_path):
    from docx import Document

    document = Document()
    document.add_heading("Compte rendu", level=1)
    document.add_paragraph("Réunion du lundi.")
    tableau = document.add_table(rows=1, cols=2)
    tableau.cell(0, 0).text, tableau.cell(0, 1).text = "Poste", "Total"
    chemin = tmp_path / "cr.docx"
    document.save(chemin)

    assert _texte(connecteur, chemin) == "Compte rendu\nRéunion du lundi.\nPoste | Total\n"


def test_un_classeur_donne_une_partie_par_feuille(connecteur, tmp_path):
    from openpyxl import Workbook

    classeur = Workbook()
    classeur.active.title = "Janvier"
    classeur.active.append(["Loyer", 150000])
    classeur.create_sheet("Fevrier").append(["Loyer", 160000])
    chemin = tmp_path / "budget.xlsx"
    classeur.save(chemin)

    texte = _texte(connecteur, chemin)

    assert "--- budget.xlsx (Janvier) ---\nLoyer | 150000" in texte
    assert "--- budget.xlsx (Fevrier) ---\nLoyer | 160000" in texte


def test_une_presentation_donne_une_partie_par_diapositive(connecteur, tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    for titre in ("Bilan", "Objectifs"):
        diapositive = presentation.slides.add_slide(presentation.slide_layouts[1])
        diapositive.shapes.title.text = titre
    chemin = tmp_path / "revue.pptx"
    presentation.save(chemin)

    texte = _texte(connecteur, chemin)

    assert "--- revue.pptx, page 1 ---\nBilan" in texte
    assert "--- revue.pptx, page 2 ---\nObjectifs" in texte


def test_un_document_sans_texte_echoue_au_lieu_d_ecrire_un_fichier_vide(connecteur, tmp_path):
    from docx import Document

    chemin = tmp_path / "vide.docx"
    Document().save(chemin)

    resultat = connecteur._convertir_un_fichier(str(chemin), "txt")

    assert resultat.statut.value == "FAILED"
    assert "VIDE" in resultat.message
    assert not list((tmp_path / "sorties").glob("*.txt")), "aucun fichier laisse derriere"


@pytest.mark.parametrize("source", ["pdf", "docx", "xlsx", "pptx"])
def test_les_quatre_formats_sont_declares(source):
    assert [m.moteur_id for m in moteurs_pour(source, "txt")] == ["lecteur-documents"]


def test_une_page_lue_par_ocr_le_dit_meme_seule(monkeypatch, tmp_path):
    """Un texte devine par OCR ne passe jamais pour un texte encode. La
    lecture est simulee : l'OCR reel depend de Tesseract sur la machine."""
    from core.production.conversion import extraction
    from tools.documents.reader import Document, Passage

    chemin = tmp_path / "scan.pdf"
    monkeypatch.setattr(extraction, "lire_document", lambda _: Document(
        chemin=chemin, statut="LU",
        passages=[Passage(texte="Devis 2026-118", fichier="scan.pdf", page=1, via_ocr=True)]))

    extraction.document_vers_texte(chemin, tmp_path / "scan.txt")

    assert (tmp_path / "scan.txt").read_text(encoding="utf-8") == \
        "--- scan.pdf, page 1 (OCR) ---\nDevis 2026-118\n"


# --- PDF -> Word sans LibreOffice (DEC-0172) --------------------------------------

def _pdf_deux_pages(tmp_path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    chemin = tmp_path / "contrat.pdf"
    page = canvas.Canvas(str(chemin), pagesize=A4)
    page.drawString(72, 750, "Article 1 : objet du contrat.")
    page.drawString(72, 730, "- pose de cloisons BA13")
    page.showPage()
    page.drawString(72, 750, "Article 2 : prix.")
    page.save()
    return chemin


def test_sans_libreoffice_un_pdf_devient_un_word_modifiable(tmp_path):
    """La machine sans LibreOffice (un PC Windows ordinaire) : le moteur de
    repli ecrit le texte, une ligne par paragraphe, une page par page."""
    from dataclasses import replace

    from docx import Document

    from core.production.conversion import registre

    originales = registre.MATRICE[("pdf", "docx")]
    registre.MATRICE[("pdf", "docx")] = [
        replace(e, disponible=lambda: (False, "LibreOffice absent")) if e.moteur_id == "libreoffice" else e
        for e in originales
    ]
    try:
        resultat = ConnecteurFileConversion(dossier=tmp_path / "sorties")._convertir_un_fichier(
            str(_pdf_deux_pages(tmp_path)), "docx")
    finally:
        registre.MATRICE[("pdf", "docx")] = originales

    assert resultat.statut.value == "SUCCESS", resultat.message
    assert resultat.detail["moteur"] == "lecteur-documents"
    assert "Texte seul" in resultat.detail["limites_qualite"]
    document = Document(resultat.preuve)
    textes = [p.text for p in document.paragraphs if p.text]
    assert textes == ["Article 1 : objet du contrat.", "- pose de cloisons BA13", "Article 2 : prix."]
    sauts = sum(1 for p in document.paragraphs for r in p.runs if 'w:br w:type="page"' in r._r.xml)
    assert sauts == 1, "un saut de page entre les deux pages"


def test_libreoffice_garde_la_priorite_quand_il_est_la():
    assert [m.moteur_id for m in moteurs_pour("pdf", "docx")] == ["libreoffice", "lecteur-documents"]


def test_une_page_ocr_est_signalee_dans_le_word(monkeypatch, tmp_path):
    from docx import Document

    from core.production.conversion import extraction
    from tools.documents.reader import Document as Lu
    from tools.documents.reader import Passage

    monkeypatch.setattr(extraction, "lire_document", lambda _: Lu(
        chemin=tmp_path / "scan.pdf", statut="LU",
        passages=[Passage(texte="Devis 2026-118", fichier="scan.pdf", page=1, via_ocr=True)]))

    extraction.pdf_vers_docx(tmp_path / "scan.pdf", tmp_path / "scan.docx")
    paragraphes = Document(tmp_path / "scan.docx").paragraphs

    assert "lue par reconnaissance de caracteres (OCR)" in paragraphes[0].text
    assert paragraphes[0].runs[0].italic
    assert paragraphes[1].text == "Devis 2026-118"


def test_un_pdf_illisible_ne_donne_pas_de_word(monkeypatch, tmp_path):
    from core.production.conversion import extraction
    from core.production.conversion.moteurs import MoteurEchec
    from tools.documents.reader import Document as Lu

    monkeypatch.setattr(extraction, "lire_document", lambda _: Lu(
        chemin=tmp_path / "vide.pdf", statut="VIDE", raison="aucun texte extractible"))

    with pytest.raises(MoteurEchec, match="aucun texte extractible"):
        extraction.pdf_vers_docx(tmp_path / "vide.pdf", tmp_path / "vide.docx")
    assert not (tmp_path / "vide.docx").exists()
