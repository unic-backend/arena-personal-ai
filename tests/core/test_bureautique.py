"""Excel, PowerPoint, Word ecrits pour de vrai depuis le markdown (DEC-0167).

Chaque test RELIT le fichier ecrit avec la bibliotheque du format : un
fichier qui s'ecrit mais ne se relit pas n'est pas un document.
"""
from pathlib import Path

import pytest

from apps.backend.routers.chat import format_de_document_demande
from core.connectors.file_conversion import ConnecteurFileConversion
from core.production.conversion.bureautique import en_nombre, lire_blocs
from core.production.conversion.registre import moteurs_pour

DEVIS = """# Devis BA13 — Almadies

Chantier du **lundi 5 octobre**.

## Postes

| Poste | Quantité | Prix unitaire | Total |
|---|---|---|---|
| Plaque BA13 | 40 | 4 500 | 180 000 |
| Rail 48 | 25 | 1 250,50 | 31 262,50 |

## Conditions

- Acompte de 50 %
- Livraison sous *8 jours*
1. Signature
2. Démarrage
"""


@pytest.fixture()
def connecteur(tmp_path: Path) -> ConnecteurFileConversion:
    return ConnecteurFileConversion(dossier=tmp_path)


def _ecrire(connecteur, format_cible: str) -> Path:
    resultat = connecteur._rediger(DEVIS, format_cible, "Devis Almadies")
    assert resultat.statut.value == "SUCCESS", resultat.message
    chemin = Path(resultat.preuve)
    assert chemin.suffix == f".{format_cible}" and chemin.is_file()
    return chemin


def test_un_excel_garde_le_tableau_et_ses_nombres(connecteur):
    from openpyxl import load_workbook

    classeur = load_workbook(_ecrire(connecteur, "xlsx"))
    feuille = classeur["Postes"]

    assert [c.value for c in feuille[1]] == ["Poste", "Quantité", "Prix unitaire", "Total"]
    assert [c.value for c in feuille[2]] == ["Plaque BA13", 40, 4500, 180000]
    assert feuille["C3"].value == 1250.5, "un nombre francais reste un nombre"
    assert feuille["A1"].font.bold


def test_un_excel_sans_tableau_ecrit_une_ligne_par_bloc(connecteur, tmp_path):
    from openpyxl import load_workbook

    resultat = connecteur._rediger("# Notes\n\nPremiere idee.\n\n- deuxieme", "xlsx", "Notes")
    feuille = load_workbook(resultat.preuve)["Texte"]

    assert [c.value for c in feuille["A"]] == ["Notes", "Premiere idee.", "deuxieme"]


def test_une_presentation_a_une_diapositive_par_titre(connecteur):
    from pptx import Presentation

    diapos = list(Presentation(_ecrire(connecteur, "pptx")).slides)
    titres = [d.shapes.title.text for d in diapos]

    assert titres == ["Devis BA13 — Almadies", "Postes", "Conditions"]
    puces = diapos[2].placeholders[1].text_frame.text
    assert "Acompte de 50 %" in puces and "Livraison sous 8 jours" in puces


def test_une_longue_liste_deborde_sur_une_diapositive_suite(connecteur):
    from pptx import Presentation

    texte = "## Etapes\n\n" + "\n".join(f"- etape {i}" for i in range(9))
    diapos = list(Presentation(connecteur._rediger(texte, "pptx", "Etapes").preuve).slides)

    assert [d.shapes.title.text for d in diapos] == ["Etapes", "Etapes (suite)"]


def test_un_word_garde_titres_listes_tableau_et_gras(connecteur):
    from docx import Document

    document = Document(_ecrire(connecteur, "docx"))
    styles = [(p.style.name, p.text) for p in document.paragraphs if p.text]

    assert ("Heading 1", "Devis BA13 — Almadies") in styles
    assert ("List Bullet", "Acompte de 50 %") in styles
    assert ("List Number", "Signature") in styles
    assert document.tables[0].cell(1, 0).text == "Plaque BA13"
    gras = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "lundi 5 octobre" in gras


def test_le_word_natif_passe_avant_libreoffice():
    """Il marche sans LibreOffice : c'est tout l'interet."""
    assert [m.moteur_id for m in moteurs_pour("md", "docx")] == ["python-docx"]
    assert moteurs_pour("md", "xlsx") and moteurs_pour("md", "pptx")


@pytest.mark.parametrize("phrase, attendu", [
    ("fais-moi un fichier excel du devis", "xlsx"),
    ("génère une présentation powerpoint de ce plan", "pptx"),
    ("crée un diaporama de ce projet", "pptx"),
    ("écris-moi ce compte rendu dans un fichier word", "docx"),
])
def test_ces_demandes_produisent_un_fichier_de_bureau(phrase, attendu):
    assert format_de_document_demande(phrase) == attendu


@pytest.mark.parametrize("cellule, valeur", [
    ("40", 40), ("1 250,50", 1250.5), ("12.5", 12.5), ("1.250.000", 1250000),
    ("5 000 FCFA", None), ("BA13", None), ("", None),
])
def test_un_nombre_francais_devient_un_nombre_une_unite_reste_du_texte(cellule, valeur):
    assert en_nombre(cellule) == valeur


def test_le_code_entre_backticks_ne_devient_pas_une_liste():
    blocs = lire_blocs("```\n- pas une puce\n```\n- une puce")

    assert [(b.genre, b.texte) for b in blocs] == [("paragraphe", "- pas une puce"), ("puce", "une puce")]
