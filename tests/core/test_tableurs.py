"""CSV <-> Excel (DEC-0169) : chaque test rouvre le fichier ecrit avec la
bibliotheque du format — un fichier qui existe n'est pas un fichier juste."""
import csv
from pathlib import Path

import pytest

from core.connectors.file_conversion import ConnecteurFileConversion
from core.production.conversion import tableurs
from core.production.conversion.bureautique import en_nombre
from core.production.conversion.moteurs import MoteurEchec

EXPORT_EXCEL_FR = (
    "Client;Montant;Code postal;Note\r\n"
    "Diallo;1 250,50;00221;=HYPERLINK(\"http://x.test\")\r\n"
    "Ndiaye;40;12500;RAS\r\n"
)


@pytest.fixture()
def connecteur(tmp_path: Path) -> ConnecteurFileConversion:
    return ConnecteurFileConversion(dossier=tmp_path / "sorties")


def _convertir(connecteur, source: Path, cible: str) -> Path:
    resultat = connecteur._convertir_un_fichier(str(source), cible)
    assert resultat.statut.value == "SUCCESS", resultat.message
    assert resultat.detail["moteur"] == "openpyxl"
    return Path(resultat.preuve)


def _ecrire(tmp_path: Path, nom: str, contenu: str, encodage: str = "utf-8") -> Path:
    chemin = tmp_path / nom
    chemin.write_bytes(contenu.encode(encodage))
    return chemin


# --- CSV -> Excel ------------------------------------------------------------------

def test_un_export_excel_francais_devient_un_vrai_classeur(connecteur, tmp_path):
    from openpyxl import load_workbook

    classeur = load_workbook(_convertir(connecteur, _ecrire(tmp_path, "clients.csv", EXPORT_EXCEL_FR), "xlsx"))
    feuille = classeur.active

    assert feuille.title == "clients"
    assert [c.value for c in feuille[1]] == ["Client", "Montant", "Code postal", "Note"]
    assert feuille["A1"].font.bold and feuille.freeze_panes == "A2"
    assert feuille["B2"].value == 1250.5, "un nombre francais devient un nombre"
    assert feuille["B3"].value == 40
    assert feuille["C2"].value == "00221", "un code a zero en tete garde ses zeros"


def test_une_formule_dans_le_csv_reste_un_texte(connecteur, tmp_path):
    from openpyxl import load_workbook

    feuille = load_workbook(
        _convertir(connecteur, _ecrire(tmp_path, "piege.csv", EXPORT_EXCEL_FR), "xlsx")).active

    assert feuille["D2"].data_type == "s"
    assert feuille["D2"].value == '=HYPERLINK("http://x.test")'


@pytest.mark.parametrize("contenu, separateur", [
    ("a,b,c\n1,2,3\n", ","),
    ("a;b;c\n1,5;2;3\n", ";"),
    ("a\tb\tc\n1\t2\t3\n", "\t"),
    ("a|b|c\n1|2|3\n", "|"),
])
def test_le_separateur_est_reconnu(contenu, separateur):
    assert tableurs.separateur(contenu) == separateur


def test_un_vieil_export_windows_1252_garde_ses_accents(connecteur, tmp_path):
    from openpyxl import load_workbook

    source = _ecrire(tmp_path, "vieux.csv", "Ville;Région\r\nThiès;Thiès\r\n", "cp1252")
    feuille = load_workbook(_convertir(connecteur, source, "xlsx")).active

    assert [c.value for c in feuille[2]] == ["Thiès", "Thiès"]
    assert feuille["B1"].value == "Région"


def test_un_encodage_illisible_echoue_en_le_disant(connecteur, tmp_path):
    """0x81 n'existe ni en UTF-8 seul ni en Windows-1252 : rien n'est devine."""
    source = tmp_path / "casse.csv"
    source.write_bytes(b"Nom;Ville\r\nDiallo;\x81\x8d\r\n")

    resultat = connecteur._convertir_un_fichier(str(source), "xlsx")

    assert resultat.statut.value == "FAILED"
    assert "encodage du CSV illisible" in resultat.message


def test_un_csv_sans_en_tete_n_a_pas_de_ligne_en_gras(connecteur, tmp_path):
    from openpyxl import load_workbook

    feuille = load_workbook(
        _convertir(connecteur, _ecrire(tmp_path, "brut.csv", "Diallo,12\nNdiaye,40\n"), "xlsx")).active

    assert feuille["B1"].value == 12
    assert not feuille["A1"].font.bold and feuille.freeze_panes is None


def test_un_csv_vide_echoue_au_lieu_d_ecrire_un_classeur_vide(connecteur, tmp_path):
    resultat = connecteur._convertir_un_fichier(str(_ecrire(tmp_path, "vide.csv", "\n ; \n")), "xlsx")

    assert resultat.statut.value == "FAILED"
    assert "aucune ligne" in resultat.message


def test_au_dela_de_ce_qu_excel_tient_le_moteur_refuse(monkeypatch, tmp_path):
    monkeypatch.setattr(tableurs, "LIGNES_MAX", 2)
    with pytest.raises(MoteurEchec, match="lignes"):
        tableurs.csv_vers_xlsx(_ecrire(tmp_path, "long.csv", "a\nb\nc\n"), tmp_path / "long.xlsx")


# --- Excel -> CSV ------------------------------------------------------------------

def _classeur(tmp_path: Path) -> Path:
    from datetime import date

    from openpyxl import Workbook

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Ventes"
    feuille.append(["Client", "Montant", "Date", "Note"])
    feuille.append(["Diallo", 1250.5, date(2026, 9, 29), "-2 % de remise"])
    feuille.append(["Ndiaye", 40, None, "@commercial"])
    autre = classeur.create_sheet("Brouillon")
    autre.append(["ne doit pas sortir"])
    chemin = tmp_path / "ventes.xlsx"
    classeur.save(chemin)
    return chemin


def test_un_classeur_devient_un_csv_qu_un_excel_francais_rouvre(connecteur, tmp_path):
    sortie = _convertir(connecteur, _classeur(tmp_path), "csv")

    brut = sortie.read_bytes()
    assert brut.startswith(b"\xef\xbb\xbf"), "BOM UTF-8 : Excel lit les accents"
    lignes = list(csv.reader(brut.decode("utf-8-sig").splitlines(), delimiter=";"))
    assert lignes[0] == ["Client", "Montant", "Date", "Note"]
    assert lignes[1][:3] == ["Diallo", "1250,5", "29/09/2026"]
    assert lignes[2][:2] == ["Ndiaye", "40"]
    assert "ne doit pas sortir" not in brut.decode("utf-8-sig"), "seule la feuille active"


def test_un_texte_qui_ressemble_a_une_formule_est_neutralise(connecteur, tmp_path):
    lignes = list(csv.reader(
        _convertir(connecteur, _classeur(tmp_path), "csv").read_text(encoding="utf-8-sig").splitlines(),
        delimiter=";"))

    assert lignes[1][3] == "'-2 % de remise"
    assert lignes[2][3] == "'@commercial"


def test_l_aller_retour_garde_les_valeurs(connecteur, tmp_path):
    from openpyxl import load_workbook

    csv_1 = _convertir(connecteur, _ecrire(tmp_path, "aller.csv", "Poste;Total\nBA13;180 000\nRail;31 262,50\n"), "xlsx")
    retour = _convertir(ConnecteurFileConversion(dossier=tmp_path / "retour"), csv_1, "csv")
    feuille = load_workbook(_convertir(ConnecteurFileConversion(dossier=tmp_path / "re"), retour, "xlsx")).active

    assert [c.value for c in feuille[2]] == ["BA13", 180000]
    assert [c.value for c in feuille[3]] == ["Rail", 31262.5]


# --- Les nombres -------------------------------------------------------------------

@pytest.mark.parametrize("cellule, valeur", [
    ("0,5", 0.5), ("0", 0), ("-3", -3),
    ("00221", None), ("0612", None),                # des codes, pas des nombres
    ("1234567890123456789", None),                  # un tableur l'arrondirait
])
def test_un_code_n_est_pas_un_nombre(cellule, valeur):
    assert en_nombre(cellule) == valeur


def test_le_markdown_vers_excel_n_ecrit_plus_de_formule(tmp_path):
    """Defaut du moteur livre le matin meme (DEC-0167), trouve en ecrivant
    celui-ci : openpyxl faisait d'une case « =... » une formule."""
    from openpyxl import load_workbook

    from core.production.conversion.bureautique import markdown_vers_xlsx

    source = _ecrire(tmp_path, "t.md", "| Nom | Lien |\n|---|---|\n| x | =WEBSERVICE(\"http://x.test\") |\n")
    markdown_vers_xlsx(source, tmp_path / "t.xlsx")
    case = load_workbook(tmp_path / "t.xlsx").active["B2"]

    assert case.data_type == "s" and case.value.startswith("=WEBSERVICE")


# --- Markdown -> CSV (DEC-0171) ----------------------------------------------------

REPONSE_AVEC_TABLEAUX = """# Ventes de septembre

| Client | Montant | Note |
|---|---|---|
| **Diallo** | 1 250,50 | =HYPERLINK("http://x.test") |
| Ndiaye | 40 | code 00221 |

## Autre tableau

| Ne | doit pas sortir |
|---|---|
| x | y |
"""


def test_une_reponse_devient_le_csv_de_son_premier_tableau(connecteur):
    resultat = connecteur._rediger(REPONSE_AVEC_TABLEAUX, "csv", "Ventes")
    assert resultat.statut.value == "SUCCESS", resultat.message
    brut = Path(resultat.preuve).read_bytes()

    assert Path(resultat.preuve).suffix == ".csv" and brut.startswith(b"\xef\xbb\xbf")
    lignes = list(csv.reader(brut.decode("utf-8-sig").splitlines(), delimiter=";"))
    assert lignes == [
        ["Client", "Montant", "Note"],
        ["Diallo", "1250,5", "'=HYPERLINK(\"http://x.test\")"],
        ["Ndiaye", "40", "code 00221"],
    ]


def test_une_reponse_sans_tableau_ne_donne_pas_de_csv(connecteur):
    resultat = connecteur._rediger("Pas de tableau ici, juste une phrase.", "csv", "Rien")

    assert resultat.statut.value == "FAILED"
    assert "aucun tableau" in resultat.message


@pytest.mark.parametrize("phrase", [
    "fais-moi un csv de ce tableau", "exporte les ventes en csv", "génère un fichier csv",
])
def test_ces_demandes_produisent_un_csv(phrase):
    from apps.backend.routers.chat import format_de_document_demande

    assert format_de_document_demande(phrase) == "csv"


def test_lire_un_csv_fourni_n_en_fabrique_pas_un():
    from apps.backend.routers.chat import format_de_document_demande

    assert format_de_document_demande("lis ce fichier csv") is None
