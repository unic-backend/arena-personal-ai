"""PowerPoint -> PDF sans LibreOffice (DEC-0175) : le PDF ecrit est relu,
une page par diapositive, et rien de la presentation n'est interprete."""
from dataclasses import replace
from pathlib import Path

import pytest

from core.connectors.file_conversion import ConnecteurFileConversion
from core.production.conversion import presentation_pdf, registre
from core.production.conversion.moteurs import MoteurEchec


def _presentation(tmp_path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)  # 16:9

    bilan = presentation.slides.add_slide(presentation.slide_layouts[1])
    bilan.shapes.title.text = "Bilan <T3>"
    cadre = bilan.placeholders[1].text_frame
    cadre.text = "Chantier Almadies"
    detail = cadre.add_paragraph()
    detail.text, detail.level = "cloisons posees", 1
    gras = cadre.add_paragraph().add_run()
    gras.text, gras.font.bold = "urgent", True

    presentation.slides.add_slide(presentation.slide_layouts[6])  # vierge : une image, par exemple

    chiffres = presentation.slides.add_slide(presentation.slide_layouts[5])
    chiffres.shapes.title.text = "Chiffres"
    tableau = chiffres.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tableau.cell(0, 0).text, tableau.cell(0, 1).text = "Poste", "Total"
    tableau.cell(1, 0).text, tableau.cell(1, 1).text = "BA13", "<script>x</script>"

    chemin = tmp_path / "revue.pptx"
    presentation.save(chemin)
    return chemin


def test_chaque_diapositive_devient_une_section(tmp_path):
    page = presentation_pdf.pptx_vers_html(_presentation(tmp_path))

    assert page.count("<section>") == 3
    assert "<h1>Bilan &lt;T3&gt;</h1>" in page
    assert '<p style="margin-left:18pt">cloisons posees</p>' in page
    assert "<strong>urgent</strong>" in page
    assert "(diapositive sans texte)" in page, "la diapositive vide garde sa page"
    assert "<td>&lt;script&gt;x&lt;/script&gt;</td>" in page and "<script>" not in page
    assert "size: 339mm 190mm" in page, "le format 16:9 de la presentation"


def test_sans_libreoffice_une_presentation_devient_un_pdf(tmp_path):
    from pypdf import PdfReader

    originales = registre.MATRICE[("pptx", "pdf")]
    registre.MATRICE[("pptx", "pdf")] = [
        replace(e, disponible=lambda: (False, "LibreOffice absent")) if e.moteur_id == "libreoffice" else e
        for e in originales
    ]
    try:
        resultat = ConnecteurFileConversion(dossier=tmp_path / "sorties")._convertir_un_fichier(
            str(_presentation(tmp_path)), "pdf")
    finally:
        registre.MATRICE[("pptx", "pdf")] = originales

    assert resultat.statut.value == "SUCCESS", resultat.message
    assert resultat.detail["moteur"] == "python-pptx+weasyprint"
    assert "notes de l'orateur" in resultat.detail["limites_qualite"]
    pages = PdfReader(resultat.preuve).pages
    assert len(pages) == 3, "une page par diapositive"
    assert "Bilan <T3>" in pages[0].extract_text()
    assert "diapositive sans texte" in pages[1].extract_text()
    assert "BA13" in pages[2].extract_text()
    assert pages[0].mediabox.width > pages[0].mediabox.height


def test_libreoffice_garde_la_priorite_pour_la_presentation():
    assert [m.moteur_id for m in registre.moteurs_pour("pptx", "pdf")] == [
        "libreoffice", "python-pptx+weasyprint"]


def test_une_presentation_sans_texte_ou_corrompue_echoue(tmp_path):
    from pptx import Presentation

    vide = Presentation()
    vide.slides.add_slide(vide.slide_layouts[6])
    chemin = tmp_path / "vide.pptx"
    vide.save(chemin)
    with pytest.raises(MoteurEchec, match="aucun texte"):
        presentation_pdf.pptx_vers_html(chemin)

    casse = tmp_path / "casse.pptx"
    casse.write_bytes(b"PK\x03\x04 pas une presentation")
    with pytest.raises(MoteurEchec, match="illisible"):
        presentation_pdf.pptx_vers_html(casse)
