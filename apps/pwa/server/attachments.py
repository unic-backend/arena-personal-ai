"""Validated, temporary attachment storage and extraction for Usman."""

from __future__ import annotations

import base64
import io
import os
import re
import tempfile
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from mutagen import File as MutagenFile
from PIL import Image
from pypdf import PdfReader


MAX_FILE_BYTES = int(os.getenv("USMAN_MAX_FILE_MB", "25")) * 1024 * 1024
MAX_EXTRACTED_CHARS = int(os.getenv("USMAN_MAX_EXTRACTED_CHARS", "120000"))
FILE_TTL_SECONDS = int(os.getenv("USMAN_FILE_TTL_SECONDS", "3600"))
ROOT = Path(tempfile.gettempdir()) / "usman-attachments"
ROOT.mkdir(mode=0o700, parents=True, exist_ok=True)

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".xml", ".html", ".htm",
    ".yaml", ".yml", ".log", ".js", ".jsx", ".ts", ".tsx", ".py", ".css",
    ".sql",
}


class AttachmentError(ValueError):
    pass


@dataclass
class StoredAttachment:
    id: str
    name: str
    size: int
    mime_type: str
    kind: str
    path: Path
    created_at: float
    metadata: dict = field(default_factory=dict)
    extracted_text: str = ""

    def public(self) -> dict:
        return {
            "id": self.id,
            "name": self.name,
            "size": self.size,
            "type": self.mime_type,
            "kind": self.kind,
            "metadata": self.metadata,
            "extractedCharacters": len(self.extracted_text),
        }


STORE: dict[str, StoredAttachment] = {}


def safe_name(name: str) -> str:
    clean = Path(name or "attachment").name
    clean = re.sub(r"[^a-zA-Z0-9._ -]", "_", clean).strip(" .")
    return (clean or "attachment")[:160]


def detect_kind(name: str, mime_type: str) -> str:
    suffix = Path(name).suffix.lower()
    mime = (mime_type or "").lower()
    if mime.startswith("image/"):
        return "image"
    if mime == "application/pdf" or suffix == ".pdf":
        return "pdf"
    if mime.startswith("audio/"):
        return "audio"
    if mime.startswith("video/"):
        return "video"
    if suffix == ".docx" or "wordprocessingml" in mime:
        return "document"
    if mime.startswith("text/") or suffix in TEXT_EXTENSIONS or "json" in mime or "xml" in mime:
        return "document"
    raise AttachmentError("unsupported attachment type")


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)[:MAX_EXTRACTED_CHARS]
        except UnicodeDecodeError:
            continue
    raise AttachmentError("document text encoding is unsupported")


def inspect_file(name: str, mime_type: str, kind: str, data: bytes) -> tuple[dict, str]:
    suffix = Path(name).suffix.lower()
    metadata: dict = {"format": suffix.lstrip(".").upper() or mime_type}
    text = ""

    if kind == "image":
        try:
            with Image.open(io.BytesIO(data)) as image:
                image.verify()
            with Image.open(io.BytesIO(data)) as image:
                metadata.update({"width": image.width, "height": image.height, "format": image.format})
        except Exception as exc:
            raise AttachmentError(f"invalid image: {exc}") from exc
    elif kind == "pdf":
        if not data.startswith(b"%PDF-"):
            raise AttachmentError("invalid PDF signature")
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = []
            remaining = MAX_EXTRACTED_CHARS
            for page in reader.pages:
                if remaining <= 0:
                    break
                value = (page.extract_text() or "")[:remaining]
                pages.append(value)
                remaining -= len(value)
            text = "\n\n".join(pages)
            metadata.update({"pages": len(reader.pages), "encrypted": reader.is_encrypted})
        except Exception as exc:
            raise AttachmentError(f"PDF extraction failed: {exc}") from exc
    elif kind == "document" and suffix == ".docx":
        try:
            document = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs)[:MAX_EXTRACTED_CHARS]
            metadata.update({"paragraphs": len(document.paragraphs), "format": "DOCX"})
        except Exception as exc:
            raise AttachmentError(f"DOCX extraction failed: {exc}") from exc
    elif kind == "document":
        text = decode_text(data)
        metadata.update({"characters": len(text), "lines": text.count("\n") + (1 if text else 0)})
    elif kind in {"audio", "video"}:
        try:
            media = MutagenFile(io.BytesIO(data))
            info = getattr(media, "info", None)
            if info:
                length = getattr(info, "length", None)
                bitrate = getattr(info, "bitrate", None)
                if length is not None:
                    metadata["duration"] = round(float(length), 3)
                if bitrate is not None:
                    metadata["bitrate"] = int(bitrate)
        except Exception:
            # The browser already probed media metadata. Lack of a Mutagen parser is not fabrication.
            metadata["metadataNote"] = "container accepted; detailed codec metadata unavailable"

    return metadata, text


def purge_expired() -> None:
    cutoff = time.time() - FILE_TTL_SECONDS
    for attachment_id, value in list(STORE.items()):
        if value.created_at >= cutoff:
            continue
        try:
            value.path.unlink(missing_ok=True)
        finally:
            STORE.pop(attachment_id, None)


def store_attachment(name: str, mime_type: str, claimed_kind: str, data: bytes) -> StoredAttachment:
    purge_expired()
    if not data:
        raise AttachmentError("empty file")
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(f"file exceeds {MAX_FILE_BYTES // 1024 // 1024} MB")
    name = safe_name(name)
    kind = detect_kind(name, mime_type)
    if claimed_kind and claimed_kind != kind:
        raise AttachmentError(f"file type mismatch: claimed {claimed_kind}, detected {kind}")
    metadata, text = inspect_file(name, mime_type, kind, data)
    attachment_id = f"file_{uuid.uuid4().hex}"
    path = ROOT / attachment_id
    path.write_bytes(data)
    try:
        path.chmod(0o600)
    except OSError:
        pass
    stored = StoredAttachment(
        id=attachment_id,
        name=name,
        size=len(data),
        mime_type=mime_type or "application/octet-stream",
        kind=kind,
        path=path,
        created_at=time.time(),
        metadata=metadata,
        extracted_text=text,
    )
    STORE[attachment_id] = stored
    return stored


def get_attachments(ids: list[str]) -> list[StoredAttachment]:
    purge_expired()
    result = []
    for attachment_id in ids[:5]:
        value = STORE.get(str(attachment_id))
        if not value:
            raise AttachmentError(f"attachment not found or expired: {attachment_id}")
        result.append(value)
    return result


def data_base64(value: StoredAttachment) -> str:
    return base64.b64encode(value.path.read_bytes()).decode("ascii")


def provider_image(value: StoredAttachment) -> tuple[str, str]:
    """Normalize a potentially huge image for fast, broadly compatible vision input."""
    with Image.open(value.path) as image:
        image.thumbnail((1800, 1800))
        if image.mode not in {"RGB", "L"}:
            background = Image.new("RGB", image.size, "white")
            if "A" in image.getbands():
                background.paste(image, mask=image.getchannel("A"))
            else:
                background.paste(image)
            image = background
        elif image.mode == "L":
            image = image.convert("RGB")
        output = io.BytesIO()
        image.save(output, format="JPEG", quality=86, optimize=True)
    return "image/jpeg", base64.b64encode(output.getvalue()).decode("ascii")


def attachment_text_context(values: list[StoredAttachment]) -> str:
    blocks = []
    remaining = MAX_EXTRACTED_CHARS
    for value in values:
        details = ", ".join(f"{key}={val}" for key, val in value.metadata.items())
        if value.extracted_text:
            excerpt = value.extracted_text[:remaining]
            remaining -= len(excerpt)
            blocks.append(
                f"<attachment name={json_quote(value.name)} type={json_quote(value.mime_type)}>\n"
                f"{excerpt}\n</attachment>"
            )
        else:
            blocks.append(
                f"<attachment name={json_quote(value.name)} type={json_quote(value.mime_type)} "
                f"metadata={json_quote(details)} />"
            )
        if remaining <= 0:
            break
    if not blocks:
        return ""
    return "\n\nAttached file contents and metadata:\n" + "\n\n".join(blocks)


def json_quote(value: str) -> str:
    return '"' + value.replace("&", "&amp;").replace('"', "&quot;") + '"'
