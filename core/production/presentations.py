"""Génération native de présentations Arena, sans code Dashi."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

THEMES = {
    "clair": {"fond": "F7F7F5", "texte": "18181B", "accent": "2563EB"},
    "sombre": {"fond": "111827", "texte": "F9FAFB", "accent": "60A5FA"},
    "minimal": {"fond": "FFFFFF", "texte": "111111", "accent": "111111"},
}
MAX_SLIDES = 40
MAX_PUCES = 12


class PresentationInvalide(ValueError):
    pass


def _texte(value: Any, max_chars: int) -> str:
    return str(value or "").strip()[:max_chars]


def normaliser_plan(plan: Dict[str, Any]) -> Dict[str, Any]:
    if not isinstance(plan, dict):
        raise PresentationInvalide("Le plan doit être un objet.")
    titre = _texte(plan.get("titre"), 180)
    if not titre:
        raise PresentationInvalide("Le titre de la présentation est requis.")
    theme = _texte(plan.get("theme") or "clair", 30).lower()
    if theme not in THEMES:
        raise PresentationInvalide(f"Thème inconnu : {theme}.")
    slides = plan.get("slides")
    if not isinstance(slides, list) or not slides:
        raise PresentationInvalide("Au moins une slide est requise.")
    if len(slides) > MAX_SLIDES:
        raise PresentationInvalide(f"Maximum {MAX_SLIDES} slides.")
    propres: List[Dict[str, Any]] = []
    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise PresentationInvalide(f"Slide {index} invalide.")
        stitre = _texte(slide.get("titre"), 180)
        if not stitre:
            raise PresentationInvalide(f"Slide {index} sans titre.")
        puces_brutes = slide.get("puces") or []
        if not isinstance(puces_brutes, list):
            raise PresentationInvalide(f"Les puces de la slide {index} doivent être une liste.")
        puces = [_texte(p, 500) for p in puces_brutes[:MAX_PUCES] if _texte(p, 500)]
        propres.append({"titre": stitre, "puces": puces})
    return {"titre": titre, "theme": theme, "slides": propres}


def generer_pptx(plan: Dict[str, Any], sortie: Path) -> Dict[str, Any]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    plan = normaliser_plan(plan)
    palette = THEMES[plan["theme"]]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def couleur(hexstr: str) -> RGBColor:
        return RGBColor.from_string(hexstr)
    for numero, spec in enumerate(plan["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fond = slide.background.fill
        fond.solid()
        fond.fore_color.rgb = couleur(palette["fond"])
        titre_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(1.1))
        titre_p = titre_box.text_frame.paragraphs[0]
        titre_p.text = spec["titre"]
        titre_p.font.size = Pt(30 if numero else 36)
        titre_p.font.bold = True
        titre_p.font.color.rgb = couleur(palette["texte"])
        if spec["puces"]:
            body = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.1), Inches(4.5))
            tf = body.text_frame
            tf.clear()
            for i, item in enumerate(spec["puces"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = item
                p.font.size = Pt(20)
                p.font.color.rgb = couleur(palette["texte"])
                p.space_after = Pt(10)
        accent = slide.shapes.add_shape(1, Inches(0.8), Inches(1.65), Inches(1.2), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = couleur(palette["accent"])
        accent.line.fill.background()
    sortie.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(sortie))
    relu = Presentation(str(sortie))
    if len(relu.slides) != len(plan["slides"]):
        raise PresentationInvalide("Le PPTX écrit ne contient pas le nombre de slides attendu.")
    return {"slides": len(relu.slides), "theme": plan["theme"], "titre": plan["titre"]}
