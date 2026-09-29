"""Markdown -> Excel, PowerPoint, Word : les fichiers de bureau que JARVIS livre.

La consigne JARVIS du proprietaire (DEC-0163, `config/jarvis.md`) : « make me
an Excel file », « create a presentation », « create a Word document » — et
« the result must be an actual file, not merely text pretending to be a
file ». Jusqu'au 29/09/2026, `rediger` repondait honnetement « aucun moteur
ARENA n'ecrit un document .xlsx / .pptx », alors que `openpyxl` et
`python-pptx` etaient deja installes (pour LIRE ces fichiers,
`tools/documents/reader.py`). Le Word, lui, dependait de LibreOffice.

Ces trois moteurs ecrivent a partir du markdown que le modele produit
spontanement, sans binaire systeme :

- **Excel** : chaque tableau markdown devient une feuille ; les nombres
  (« 1 250,50 », « 12.5 ») deviennent des nombres, pas du texte. Un texte
  sans tableau donne une feuille d'une colonne, une ligne par ligne — et la
  limite de qualite le dit.
- **PowerPoint** : chaque titre (`#`, `##`) ouvre une diapositive ; ses
  paragraphes et listes deviennent ses puces. Sans titre, le texte est
  decoupe en diapositives de six puces.
- **Word** : titres, paragraphes, listes a puces et numerotees, tableaux,
  gras et italique.

Ce qu'ils ne font pas, et que `limites_qualite` dit : pas de graphique, pas
d'image, pas de theme ; le markdown exotique (citations imbriquees, HTML
brut) passe en texte.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Tuple

from core.production.conversion.moteurs import MoteurEchec

# --- Le markdown, lu en blocs ----------------------------------------------------


@dataclass
class Bloc:
    """Un morceau du document : `titre`, `paragraphe`, `puce`, `numero`, `tableau`."""

    genre: str
    texte: str = ""
    niveau: int = 0
    lignes: List[List[str]] = field(default_factory=list)


_TITRE = re.compile(r"^(#{1,6})\s+(.*)$")
_PUCE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMERO = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_SEPARATEUR_DE_TABLEAU = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cellules(ligne: str) -> List[str]:
    return [c.strip() for c in ligne.strip().strip("|").split("|")]


def lire_blocs(texte: str) -> List[Bloc]:
    """Le markdown en blocs, dans l'ordre. Le code entre ``` passe en paragraphe."""
    blocs: List[Bloc] = []
    lignes = (texte or "").splitlines()
    i = 0
    paragraphe: List[str] = []

    def fermer_paragraphe() -> None:
        if paragraphe:
            blocs.append(Bloc("paragraphe", " ".join(p.strip() for p in paragraphe)))
            paragraphe.clear()

    while i < len(lignes):
        ligne = lignes[i]
        nue = ligne.strip()
        if not nue:
            fermer_paragraphe()
            i += 1
            continue
        if nue.startswith("```"):
            fermer_paragraphe()
            code = []
            i += 1
            while i < len(lignes) and not lignes[i].strip().startswith("```"):
                code.append(lignes[i])
                i += 1
            blocs.append(Bloc("paragraphe", "\n".join(code)))
            i += 1
            continue
        titre = _TITRE.match(nue)
        if titre:
            fermer_paragraphe()
            blocs.append(Bloc("titre", titre.group(2).strip(), niveau=len(titre.group(1))))
            i += 1
            continue
        if "|" in nue and i + 1 < len(lignes) and _SEPARATEUR_DE_TABLEAU.match(lignes[i + 1]):
            fermer_paragraphe()
            tableau = [_cellules(nue)]
            i += 2
            while i < len(lignes) and "|" in lignes[i] and lignes[i].strip():
                tableau.append(_cellules(lignes[i]))
                i += 1
            blocs.append(Bloc("tableau", lignes=tableau))
            continue
        puce = _PUCE.match(ligne)
        if puce:
            fermer_paragraphe()
            blocs.append(Bloc("puce", puce.group(1).strip()))
            i += 1
            continue
        numero = _NUMERO.match(ligne)
        if numero:
            fermer_paragraphe()
            blocs.append(Bloc("numero", numero.group(1).strip()))
            i += 1
            continue
        if re.fullmatch(r"[-*_]{3,}", nue):
            fermer_paragraphe()
            i += 1
            continue
        paragraphe.append(nue)
        i += 1
    fermer_paragraphe()
    return blocs


_EMPHASE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__|_[^_]+_|`[^`]+`)")


def segments(texte: str) -> List[Tuple[str, bool, bool]]:
    """(texte, gras, italique) : l'emphase markdown d'une ligne, en morceaux."""
    morceaux: List[Tuple[str, bool, bool]] = []
    for partie in _EMPHASE.split(texte):
        if not partie:
            continue
        if partie.startswith(("**", "__")) and partie.endswith(("**", "__")) and len(partie) > 4:
            morceaux.append((partie[2:-2], True, False))
        elif partie.startswith("`") and partie.endswith("`"):
            morceaux.append((partie[1:-1], False, False))
        elif partie[0] in "*_" and partie[-1] == partie[0] and len(partie) > 2:
            morceaux.append((partie[1:-1], False, True))
        else:
            morceaux.append((partie, False, False))
    return morceaux


def sans_emphase(texte: str) -> str:
    return "".join(morceau for morceau, _, _ in segments(texte))


# --- Les nombres ----------------------------------------------------------------

_NOMBRE = re.compile(r"^[-+]?\d{1,3}(?:[   .]\d{3})*(?:[.,]\d+)?$|^[-+]?\d+(?:[.,]\d+)?$")


def en_nombre(cellule: str) -> Optional[float]:
    """« 1 250,50 » -> 1250.5 ; « 12.5 » -> 12.5 ; « 5 000 FCFA » -> None
    (une unite fait une cellule de texte : la convertir perdrait l'unite)."""
    brut = sans_emphase(cellule).strip()
    if not brut or not _NOMBRE.match(brut):
        return None
    compact = re.sub(r"[   ]", "", brut)
    if "," in compact:
        # Ecriture francaise : le point (s'il y en a) separe les milliers.
        compact = compact.replace(".", "").replace(",", ".")
    elif compact.count(".") > 1:
        compact = compact.replace(".", "")  # « 1.250.000 »
    try:
        valeur = float(compact)
    except ValueError:
        return None
    return int(valeur) if "." not in compact else valeur


# --- Excel ------------------------------------------------------------------------

def openpyxl_disponible() -> Tuple[bool, str]:
    try:
        import openpyxl  # noqa: F401
    except ImportError as erreur:
        return False, f"openpyxl n'est pas installé : {erreur}"
    return True, "openpyxl installé"


def _nom_de_feuille(base: str, pris: set) -> str:
    nom = re.sub(r"[\\/*?:\[\]]", " ", base).strip()[:31] or "Feuille"
    candidat, n = nom, 2
    while candidat in pris:
        suffixe = f" ({n})"
        candidat = nom[:31 - len(suffixe)] + suffixe
        n += 1
    pris.add(candidat)
    return candidat


def markdown_vers_xlsx(entree: Path, sortie: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    blocs = lire_blocs(entree.read_text(encoding="utf-8", errors="replace"))
    classeur = Workbook()
    classeur.remove(classeur.active)
    pris: set = set()
    dernier_titre = ""
    tableaux = 0
    for bloc in blocs:
        if bloc.genre == "titre":
            dernier_titre = sans_emphase(bloc.texte)
        if bloc.genre != "tableau" or not bloc.lignes:
            continue
        tableaux += 1
        feuille = classeur.create_sheet(_nom_de_feuille(dernier_titre or f"Tableau {tableaux}", pris))
        for numero, ligne in enumerate(bloc.lignes, start=1):
            for colonne, cellule in enumerate(ligne, start=1):
                nombre = None if numero == 1 else en_nombre(cellule)
                case = feuille.cell(row=numero, column=colonne,
                                    value=nombre if nombre is not None else sans_emphase(cellule))
                if numero == 1:
                    case.font = Font(bold=True)
        for colonne in feuille.columns:
            largeur = max(len(str(c.value or "")) for c in colonne)
            feuille.column_dimensions[colonne[0].column_letter].width = min(max(10, largeur + 2), 60)

    if tableaux == 0:
        # Pas de tableau : une ligne par bloc, pour que le fichier existe et
        # se relise — la limite de qualite du moteur le dit.
        feuille = classeur.create_sheet("Texte")
        rangee = 1
        for bloc in blocs:
            valeur = sans_emphase(bloc.texte)
            if not valeur:
                continue
            case = feuille.cell(row=rangee, column=1, value=valeur)
            if bloc.genre == "titre":
                case.font = Font(bold=True)
            rangee += 1
        if rangee == 1:
            raise MoteurEchec("aucun contenu a ecrire dans le classeur")
        feuille.column_dimensions["A"].width = 100
    classeur.save(str(sortie))


# --- PowerPoint -----------------------------------------------------------------

def python_pptx_disponible() -> Tuple[bool, str]:
    try:
        import pptx  # noqa: F401
    except ImportError as erreur:
        return False, f"python-pptx n'est pas installé : {erreur}"
    return True, "python-pptx installé"

#: Au-dela, une diapositive ne se lit plus : on en ouvre une suivante.
PUCES_PAR_DIAPOSITIVE = 6


def _diapositives(blocs: List[Bloc]) -> Tuple[str, str, List[Tuple[str, List[str]]]]:
    """(titre du document, sous-titre, [(titre de diapositive, puces)]).

    Le texte entre le titre `#` et la premiere section devient le sous-titre
    de la couverture, pas une diapositive de plus portant le meme titre."""
    titre_document = ""
    sous_titre: List[str] = []
    diapos: List[Tuple[str, List[str]]] = []
    courante: Optional[Tuple[str, List[str]]] = None
    for bloc in blocs:
        if bloc.genre == "titre":
            texte = sans_emphase(bloc.texte)
            if bloc.niveau == 1 and not titre_document and not diapos and courante is None:
                titre_document = texte
                continue
            courante = (texte, [])
            diapos.append(courante)
            continue
        if bloc.genre == "tableau":
            puces = [" — ".join(sans_emphase(c) for c in ligne) for ligne in bloc.lignes]
        else:
            puces = [sans_emphase(bloc.texte)]
        if titre_document and courante is None and bloc.genre == "paragraphe":
            sous_titre.extend(puces)
            continue
        for puce in puces:
            if courante is None or len(courante[1]) >= PUCES_PAR_DIAPOSITIVE:
                suite = courante[0] + " (suite)" if courante and courante[0] else ""
                courante = (suite.replace(" (suite) (suite)", " (suite)"), [])
                diapos.append(courante)
            courante[1].append(puce)
    return titre_document, " ".join(sous_titre), [d for d in diapos if d[0] or d[1]]


def markdown_vers_pptx(entree: Path, sortie: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    blocs = lire_blocs(entree.read_text(encoding="utf-8", errors="replace"))
    titre_document, sous_titre, diapos = _diapositives(blocs)
    if not titre_document and not diapos:
        raise MoteurEchec("aucun contenu a mettre en diapositives")
    presentation = Presentation()
    if titre_document:
        couverture = presentation.slides.add_slide(presentation.slide_layouts[0])
        couverture.shapes.title.text = titre_document
        couverture.placeholders[1].text = sous_titre
    for titre, puces in diapos:
        diapo = presentation.slides.add_slide(presentation.slide_layouts[1])
        diapo.shapes.title.text = titre or titre_document or " "
        cadre = diapo.placeholders[1].text_frame
        cadre.clear()
        for rang, puce in enumerate(puces):
            paragraphe = cadre.paragraphs[0] if rang == 0 else cadre.add_paragraph()
            paragraphe.text = puce
            paragraphe.font.size = Pt(20 if len(puces) <= 4 else 16)
    presentation.save(str(sortie))


# --- Word --------------------------------------------------------------------------

def python_docx_disponible() -> Tuple[bool, str]:
    try:
        import docx  # noqa: F401
    except ImportError as erreur:
        return False, f"python-docx n'est pas installé : {erreur}"
    return True, "python-docx installé"


def _ecrire_avec_emphase(paragraphe, texte: str) -> None:
    for morceau, gras, italique in segments(texte):
        run = paragraphe.add_run(morceau)
        run.bold = gras or None
        run.italic = italique or None


def markdown_vers_docx(entree: Path, sortie: Path) -> None:
    from docx import Document

    blocs = lire_blocs(entree.read_text(encoding="utf-8", errors="replace"))
    if not blocs:
        raise MoteurEchec("aucun contenu a ecrire dans le document")
    document = Document()
    for bloc in blocs:
        if bloc.genre == "titre":
            document.add_heading(sans_emphase(bloc.texte), level=min(bloc.niveau, 4))
        elif bloc.genre == "puce":
            _ecrire_avec_emphase(document.add_paragraph(style="List Bullet"), bloc.texte)
        elif bloc.genre == "numero":
            _ecrire_avec_emphase(document.add_paragraph(style="List Number"), bloc.texte)
        elif bloc.genre == "tableau" and bloc.lignes:
            colonnes = max(len(ligne) for ligne in bloc.lignes)
            tableau = document.add_table(rows=len(bloc.lignes), cols=colonnes)
            tableau.style = "Table Grid"
            for r, ligne in enumerate(bloc.lignes):
                for c in range(colonnes):
                    cellule = tableau.cell(r, c)
                    cellule.text = sans_emphase(ligne[c]) if c < len(ligne) else ""
                    if r == 0:
                        for run in cellule.paragraphs[0].runs:
                            run.bold = True
        else:
            _ecrire_avec_emphase(document.add_paragraph(), bloc.texte)
    document.save(str(sortie))
