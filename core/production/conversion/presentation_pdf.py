"""PowerPoint -> PDF sans LibreOffice : une page par diapositive (DEC-0175).

Dernier des couples bureautiques qui dependaient du seul LibreOffice (apres
DEC-0172 a DEC-0174). python-pptx lit chaque diapositive dans l'ordre de ses
formes : son titre, ses zones de texte (niveaux de puce compris), ses
tableaux. Le rendu passe par `word_pdf.rendre_pdf` — texte echappe, aucune
ressource chargee — sur des pages au format de la presentation (16:9 ou 4:3).

Ce que ce moteur ne reprend pas, et que la limite de qualite dit : images,
formes dessinees, couleurs et arriere-plans, disposition exacte, notes de
l'orateur. Il passe APRES LibreOffice, qui garde tout cela quand il est la.
"""
from __future__ import annotations

import html
from pathlib import Path
from typing import List

from core.production.conversion.moteurs import MoteurEchec
from core.production.conversion.word_pdf import rendre_pdf

#: Une unite EMU d'Office vaut 1/36000 de millimetre.
EMU_PAR_MM = 36000


def _style(largeur_mm: float, hauteur_mm: float) -> str:
    return f"""
@page {{ size: {largeur_mm:.0f}mm {hauteur_mm:.0f}mm; margin: 10mm 14mm; }}
body {{ font-family: 'DejaVu Sans', Arial, sans-serif; font-size: 14pt; }}
section {{ page-break-after: always; }}
section:last-child {{ page-break-after: auto; }}
h1 {{ font-size: 24pt; margin: 0 0 10pt; }}
p {{ margin: 3pt 0; }}
table {{ border-collapse: collapse; margin: 6pt 0; font-size: 11pt; }}
td {{ border: 1px solid #888; padding: 2pt 6pt; vertical-align: top; }}
"""


def _paragraphe(paragraphe) -> str:
    """Un paragraphe de zone de texte, echappe, avec gras et italique ; son
    niveau de puce devient un retrait."""
    morceaux: List[str] = []
    for run in paragraphe.runs:
        texte = html.escape(run.text)
        if not texte:
            continue
        if run.font.bold:
            texte = f"<strong>{texte}</strong>"
        if run.font.italic:
            texte = f"<em>{texte}</em>"
        morceaux.append(texte)
    contenu = "".join(morceaux)
    if not contenu.strip():
        return ""
    retrait = f' style="margin-left:{paragraphe.level * 18}pt"' if paragraphe.level else ""
    return f"<p{retrait}>{contenu}</p>"


def _tableau(tableau) -> str:
    lignes = []
    for ligne in tableau.rows:
        cases = "".join(f"<td>{html.escape(case.text)}</td>" for case in ligne.cells)
        lignes.append(f"<tr>{cases}</tr>")
    return f"<table>{''.join(lignes)}</table>"


def _diapositive(diapositive) -> str:
    """Le titre d'abord (ou qu'il soit dans l'ordre des formes), puis le reste."""
    titre = diapositive.shapes.title
    corps: List[str] = []
    if titre is not None and titre.has_text_frame and titre.text_frame.text.strip():
        corps.append(f"<h1>{html.escape(titre.text_frame.text.strip())}</h1>")
    for forme in diapositive.shapes:
        if titre is not None and forme.shape_id == titre.shape_id:
            continue
        if forme.has_text_frame:
            corps.extend(p for p in (_paragraphe(x) for x in forme.text_frame.paragraphs) if p)
        elif getattr(forme, "has_table", False) and forme.has_table:
            corps.append(_tableau(forme.table))
    return "".join(corps)


def pptx_vers_html(entree: Path) -> str:
    """Toute la presentation, une section (donc une page) par diapositive."""
    from pptx import Presentation

    try:
        presentation = Presentation(str(entree))
    except Exception as erreur:  # noqa: BLE001 — un pptx corrompu leve des types varies
        raise MoteurEchec(f"presentation illisible : {erreur}") from erreur
    sections = [_diapositive(d) for d in presentation.slides]
    if not any(sections):
        raise MoteurEchec("la presentation ne contient aucun texte")
    largeur = (presentation.slide_width or 9144000) / EMU_PAR_MM
    hauteur = (presentation.slide_height or 5143500) / EMU_PAR_MM
    # Une diapositive d'images seules garde sa page, et le dit : sans elle,
    # la numerotation du PDF ne correspondrait plus a celle de la presentation.
    pages = "".join(
        f"<section>{contenu or '<p><em>(diapositive sans texte)</em></p>'}</section>"
        for contenu in sections)
    return (f"<html><head><meta charset='utf-8'><style>{_style(largeur, hauteur)}</style>"
            f"</head><body>{pages}</body></html>")


def pptx_vers_pdf(entree: Path, sortie: Path) -> None:
    rendre_pdf(pptx_vers_html(entree), sortie)
